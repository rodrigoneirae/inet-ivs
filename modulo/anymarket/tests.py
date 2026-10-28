import json
from datetime import datetime, timedelta, timezone as dt_timezone
from io import BytesIO
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from .models import Pedido, Sincronizacion
from .services import AnyMarketError, consultar_pedidos, guardar_pedido, sincronizar


def pedido(identificador: int, estado: str, fecha: str = '2026-10-27T12:00:00Z') -> dict:
    return {'id': identificador, 'status': estado, 'createdAt': fecha, 'items': []}


class ImportacionTests(TestCase):
    def test_importa_pagados_y_actualiza_cancelados_sin_importar_pendientes(self) -> None:
        inicio = datetime(2026, 10, 27, tzinfo=dt_timezone.utc)
        guardar_pedido(pedido(1, 'PENDING'), inicio)
        guardar_pedido(pedido(2, 'PAID_WAITING_SHIP', '2026-10-26T12:00:00Z'), inicio)
        guardar_pedido(pedido(3, 'PAID_WAITING_SHIP'), inicio)
        guardar_pedido(pedido(3, 'PAID_WAITING_SHIP'), inicio)
        guardar_pedido(pedido(3, 'CANCELED'), inicio)
        self.assertEqual(Pedido.objects.count(), 1)
        self.assertEqual(Pedido.objects.get().estado, 'CANCELED')

    def test_error_no_avanza_checkpoint_ni_guarda_pedidos_parciales(self) -> None:
        ahora = datetime(2026, 10, 28, tzinfo=dt_timezone.utc)

        def consulta(_filtros: dict):
            yield pedido(1, 'PAID_WAITING_SHIP')
            raise AnyMarketError('error')

        with patch('modulo.anymarket.services.timezone.now', return_value=ahora), patch(
            'modulo.anymarket.services.consultar_pedidos', side_effect=consulta
        ), self.assertRaises(AnyMarketError):
            sincronizar()
        self.assertEqual(Pedido.objects.count(), 0)
        self.assertIsNone(Sincronizacion.objects.get().actualizado_hasta)

    def test_sincronizacion_incremental_divide_intervalos_y_actualiza_estado(self) -> None:
        ahora = datetime(2026, 11, 20, tzinfo=dt_timezone.utc)
        Sincronizacion.objects.create(clave='pedidos', actualizado_hasta=ahora - timedelta(days=14))
        Pedido.objects.create(anymarket_id=1, creado_en=ahora, estado='PAID_WAITING_SHIP', datos={})
        with patch('modulo.anymarket.services.timezone.now', return_value=ahora), patch(
            'modulo.anymarket.services.consultar_pedidos', return_value=[pedido(1, 'CANCELED')]
        ) as consulta:
            sincronizar()
        self.assertEqual(consulta.call_count, 3)
        self.assertEqual(Pedido.objects.get().estado, 'CANCELED')
        self.assertEqual(Sincronizacion.objects.get().actualizado_hasta, ahora)

    @override_settings(ANYMARKET_TOKEN='test', ANYMARKET_PLATFORM='INET', ANYMARKET_ENVIRONMENT='sandbox')
    def test_paginacion_y_cabeceras(self) -> None:
        paginas = [
            {'content': [pedido(1, 'PAID_WAITING_SHIP')], 'links': [{'rel': 'next', 'href': 'https://evil.invalid'}]},
            {'content': [pedido(2, 'PAID_WAITING_SHIP')], 'links': []},
        ]
        with patch('modulo.anymarket.services.build_opener') as opener:
            opener.return_value.open.side_effect = [BytesIO(json.dumps(p).encode()) for p in paginas]
            resultados = list(consultar_pedidos({'status': 'PAID_WAITING_SHIP'}))
        self.assertEqual(len(resultados), 2)
        request = opener.return_value.open.call_args_list[1].args[0]
        self.assertTrue(request.full_url.startswith('https://sandbox-api.anymarket.com.br/v2/orders?'))
        self.assertIn('offset=1', request.full_url)
        self.assertEqual(request.get_header('Gumgatoken'), 'test')

    @override_settings(ANYMARKET_TOKEN='test', ANYMARKET_PLATFORM='INET', ANYMARKET_ENVIRONMENT='sandbox')
    def test_ignora_enlace_next_incorrecto_en_ultima_pagina(self) -> None:
        pagina = {'content': [pedido(1, 'PAID_WAITING_SHIP')],
                  'page': {'number': 1, 'totalPages': 1}, 'links': [{'rel': 'next'}]}
        with patch('modulo.anymarket.services.build_opener') as opener:
            opener.return_value.open.return_value = BytesIO(json.dumps(pagina).encode())
            self.assertEqual(len(list(consultar_pedidos({}))), 1)
            self.assertEqual(opener.return_value.open.call_count, 1)


@override_settings(STORAGES={
    'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
    'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'},
}, ANYMARKET_TIME_ZONE='America/Santiago')
class PedidosVistaTests(TestCase):
    def setUp(self) -> None:
        self.usuario = get_user_model().objects.create_user(username='operador')
        self.url = reverse('anymarket:pedidos')

    def test_requiere_login_y_permiso(self) -> None:
        self.assertEqual(self.client.get(self.url).status_code, 302)
        self.client.force_login(self.usuario)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertNotContains(self.client.get(reverse('inicio')), 'Pedidos AnyMarket')

    def test_listado_paginado_escapa_datos_y_muestra_fecha_chilena(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        for identificador in range(26):
            Pedido.objects.create(
                anymarket_id=identificador + 1,
                creado_en=datetime(2026, 10, 27, 4, 0, tzinfo=dt_timezone.utc),
                estado='PAID_WAITING_SHIP',
                datos={'marketPlaceNumber': '<script>malicioso</script>', 'marketPlace': 'TIENDA', 'total': 0,
                       'buyer': {'name': 'Cliente', 'email': 'NO EXPONER'}},
            )
        respuesta = self.client.get(self.url, {'tamano': '25'})
        self.assertContains(respuesta, 'vendor/litepicker/litepicker.js')
        self.assertContains(respuesta, 'id="rango-fechas"')
        self.assertContains(respuesta, 'anymarket/rango-fechas.js')
        self.assertEqual(len(respuesta.context['pagina']), 25)
        self.assertContains(respuesta, '27/10/2026 01:00')
        self.assertContains(respuesta, '&lt;script&gt;malicioso&lt;/script&gt;')
        self.assertNotContains(respuesta, '<script>malicioso</script>')
        self.assertNotContains(respuesta, 'NO EXPONER')
        self.assertContains(respuesta, 'Siguiente')
        self.assertEqual(len(self.client.get(self.url, {'page': 2, 'tamano': '25'}).context['pagina']), 1)
        self.assertContains(self.client.get(reverse('inicio')), 'Pedidos AnyMarket')
        self.assertIn('no-store', respuesta['Cache-Control'])

    def test_sin_pedidos(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        self.assertContains(self.client.get(self.url), 'No hay pedidos que coincidan con los filtros.')

    @patch('modulo.anymarket.views.sincronizar_pedidos.apply_async')
    def test_sincronizacion_manual_requiere_permiso_y_limita_solicitudes(self, encolar: MagicMock) -> None:
        cache.clear()
        url = reverse('anymarket:ejecutar_sincronizacion')
        self.assertEqual(self.client.post(url).status_code, 302)
        self.client.force_login(self.usuario)
        self.assertEqual(self.client.post(url).status_code, 403)
        self.usuario.user_permissions.add(Permission.objects.get(codename='change_pedido', content_type__app_label='anymarket'))
        self.client.force_login(get_user_model().objects.get(pk=self.usuario.pk))
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertRedirects(self.client.post(url), self.url, fetch_redirect_response=False)
        self.assertRedirects(self.client.post(url), self.url, fetch_redirect_response=False)
        encolar.assert_called_once_with(retry=False)
        cache.clear()

    def test_filtros_fecha_chilena_busqueda_canal_y_rango_invalido(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        for identificador, fecha in enumerate((datetime(2026, 10, 27, 2, 59, tzinfo=dt_timezone.utc),
                                              datetime(2026, 10, 27, 3, 0, tzinfo=dt_timezone.utc)), start=1):
            Pedido.objects.create(anymarket_id=identificador, creado_en=fecha, estado='PAID_WAITING_SHIP',
                                  datos={'marketPlaceNumber': 'OC-123', 'marketPlace': 'Tienda',
                                         'buyer': {'name': 'Ana'}, 'items': [{}, {}], 'total': 100,
                                         'marketPlaceStatus': 'paid'})
        respuesta = self.client.get(self.url, {'desde': '2026-10-27', 'hasta': '2026-10-27',
                                              'busqueda': 'ana', 'canal': 'Tienda', 'estado': 'paid'})
        self.assertEqual(respuesta.context['pagina'].paginator.count, 1)
        self.assertContains(respuesta, 'Ana')
        self.assertEqual(respuesta.context['pagina'][0].cantidad_items, 2)
        self.assertContains(self.client.get(self.url, {'desde': '2026-10-28', 'hasta': '2026-10-27'}),
                            'La fecha final debe ser igual o posterior')

    def test_orden_total_across_paginas_y_parametros_invalidos(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        for numero in range(26):
            Pedido.objects.create(anymarket_id=numero + 1, creado_en=datetime(2026, 10, 27, tzinfo=dt_timezone.utc),
                                  estado='PAID_WAITING_SHIP', datos={'total': numero})
        primera = self.client.get(self.url, {'orden': 'total', 'direccion': 'desc', 'tamano': '25'})
        self.assertEqual(primera.context['pagina'][0].datos['total'], 25)
        self.assertContains(primera, 'orden=total&amp;direccion=desc&amp;tamano=25&amp;page=2')
        segunda = self.client.get(self.url, {'orden': 'total', 'direccion': 'desc', 'tamano': '25', 'page': 2})
        self.assertEqual(segunda.context['pagina'][0].datos['total'], 0)
        invalida = self.client.get(self.url, {'orden': 'datos__buyer', 'direccion': 'invalida'})
        self.assertEqual(invalida.context['orden'], 'fecha')
        self.assertEqual(invalida.context['direccion'], 'desc')

    def test_busqueda_global_abarca_columnas_y_todas_las_paginas(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        fecha = datetime(2026, 10, 27, 4, 0, tzinfo=dt_timezone.utc)
        Pedido.objects.create(anymarket_id=1, creado_en=fecha, estado='PAID_WAITING_SHIP', datos={
            'marketPlaceNumber': 'OC-ESPECIAL', 'marketPlace': 'CANAL-UNICO',
            'buyer': {'name': 'CLIENTE-UNICO'}, 'items': [{}, {}, {}],
            'total': 98765, 'marketPlaceStatus': 'ESTADO-UNICO',
        })
        for numero in range(2, 15):
            Pedido.objects.create(anymarket_id=numero, creado_en=fecha, estado='PAID_WAITING_SHIP', datos={})
        for termino in ('OC-ESPECIAL', 'CANAL-UNICO', 'CLIENTE-UNICO', 'ESTADO-UNICO', '98765', '3'):
            with self.subTest(termino=termino):
                respuesta = self.client.get(self.url, {'busqueda': termino, 'tamano': '5'})
                self.assertEqual(respuesta.context['pagina'].paginator.count, 1)
                self.assertContains(respuesta, 'OC-ESPECIAL')
        respuesta = self.client.get(self.url, {'busqueda': '27/10/2026'})
        self.assertEqual(respuesta.context['pagina'].paginator.count, 14)
        exportacion = self.client.get(reverse('anymarket:exportar_pedidos'), {'busqueda': 'CANAL-UNICO'})
        libro = load_workbook(BytesIO(b''.join(exportacion.streaming_content)), read_only=True)
        self.assertEqual(len(list(libro.active.rows)), 2)

    def test_nombres_de_canal_en_listado_busqueda_y_exportacion(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        for numero, (original, nombre) in enumerate((('mercado_livre', 'MERCADO LIBRE'),
                                                     ('SHOPIFY', 'PORTOMENAJE'),
                                                     ('C52434-09', 'YAKU')), start=1):
            with self.subTest(canal=original):
                Pedido.objects.create(anymarket_id=numero, creado_en=datetime(2026, 10, 27, tzinfo=dt_timezone.utc),
                                      estado='PAID_WAITING_SHIP', datos={'marketPlace': original})
                respuesta = self.client.get(self.url, {'busqueda': nombre, 'canal': original})
                self.assertEqual(respuesta.context['pagina'].paginator.count, 1)
                self.assertContains(respuesta, nombre)
                exportacion = self.client.get(reverse('anymarket:exportar_pedidos'), {'busqueda': nombre})
                libro = load_workbook(BytesIO(b''.join(exportacion.streaming_content)), read_only=True)
                self.assertEqual(list(libro.active.values)[1][0], nombre)

    def test_exportacion_filtrada_protegida_y_sin_formulas(self) -> None:
        url = reverse('anymarket:exportar_pedidos')
        self.assertEqual(self.client.get(url).status_code, 302)
        self.client.force_login(self.usuario)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(get_user_model().objects.get(pk=self.usuario.pk))
        for numero in range(12):
            Pedido.objects.create(anymarket_id=numero + 1, creado_en=datetime(2026, 10, 27, 4, tzinfo=dt_timezone.utc),
                                  estado='PAID_WAITING_SHIP', datos={'marketPlace': 'TIENDA', 'marketPlaceStatus': 'paid',
                                  'marketPlaceNumber': '=1+1', 'buyer': {'name': 'Cliente'}, 'total': numero})
        Pedido.objects.create(anymarket_id=100, creado_en=datetime(2026, 10, 27, tzinfo=dt_timezone.utc),
                              estado='PAID_WAITING_SHIP', datos={'marketPlace': 'OTRA', 'marketPlaceStatus': 'pending'})
        respuesta = self.client.get(url, {'canal': 'TIENDA', 'estado': 'paid', 'tamano': '5'})
        self.assertEqual(respuesta.status_code, 200)
        libro = load_workbook(BytesIO(b''.join(respuesta.streaming_content)), read_only=True)
        filas = list(libro.active.rows)
        self.assertEqual(len(filas), 13)
        self.assertEqual(filas[1][1].value, '=1+1')
        self.assertEqual(filas[1][1].data_type, 's')
        self.assertEqual(filas[1][6].value, 'paid')
        self.assertEqual(self.client.get(url, {'desde': 'invalida'}).status_code, 400)
