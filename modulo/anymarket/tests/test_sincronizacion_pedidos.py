import json
from datetime import datetime, timedelta, timezone as dt_timezone
from io import BytesIO
from threading import Event, Thread
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.cache import cache
from django.db import connections
from django.test import TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from ..models import Pedido, Sincronizacion
from ..services import AnyMarketError, consultar_pedidos, guardar_pedido, reconciliar, sincronizar


def pedido(identificador: int, estado: str, fecha: str = '2026-10-27T12:00:00Z') -> dict:
    return {'id': identificador, 'status': estado, 'createdAt': fecha, 'items': []}


class ImportacionTests(TestCase):
    def test_importa_pagados_y_actualiza_cancelados_sin_importar_pendientes(self) -> None:
        inicio = datetime(2026, 10, 27, tzinfo=dt_timezone.utc)
        guardar_pedido(pedido(1, 'PENDING'), inicio)
        guardar_pedido(pedido(2, 'PAID_WAITING_SHIP', '2026-10-26T12:00:00Z'), inicio)
        guardar_pedido(pedido(3, 'PAID_WAITING_SHIP'), inicio)
        guardar_pedido(pedido(3, 'PAID_WAITING_SHIP'), inicio)
        guardar_pedido(pedido(3, 'CANCELED'), inicio)
        self.assertEqual(Pedido.objects.count(), 1)
        self.assertEqual(Pedido.objects.get().estado, 'CANCELED')

    def test_error_no_avanza_checkpoint_ni_guarda_pedidos_parciales(self) -> None:
        ahora = datetime(2026, 10, 28, tzinfo=dt_timezone.utc)

        def consulta(_filtros: dict):
            yield pedido(1, 'PAID_WAITING_SHIP')
            raise AnyMarketError('error')

        with patch('modulo.anymarket.services.sincronizacion_pedidos.timezone.now', return_value=ahora), patch(
            'modulo.anymarket.services.sincronizacion_pedidos.consultar_pedidos', side_effect=consulta
        ), self.assertRaises(AnyMarketError):
            sincronizar()
        self.assertEqual(Pedido.objects.count(), 0)
        self.assertIsNone(Sincronizacion.objects.get().actualizado_hasta)

    def test_sincronizacion_incremental_divide_intervalos_y_actualiza_estado(self) -> None:
        ahora = datetime(2026, 11, 20, tzinfo=dt_timezone.utc)
        Sincronizacion.objects.create(clave='pedidos', actualizado_hasta=ahora - timedelta(days=14))
        Pedido.objects.create(anymarket_id=1, creado_en=ahora, estado='PAID_WAITING_SHIP', datos={})
        with patch('modulo.anymarket.services.sincronizacion_pedidos.timezone.now', return_value=ahora), patch(
            'modulo.anymarket.services.sincronizacion_pedidos.consultar_pedidos', return_value=[pedido(1, 'CANCELED')]
        ) as consulta:
            sincronizar()
        self.assertEqual(consulta.call_count, 3)
        self.assertEqual(Pedido.objects.get().estado, 'CANCELED')
        self.assertEqual(Sincronizacion.objects.get().actualizado_hasta, ahora)

    def test_lotes_confirmados_se_reprocesan_sin_duplicados_tras_error(self) -> None:
        ahora = datetime(2026, 10, 28, tzinfo=dt_timezone.utc)
        inicio = ahora - timedelta(hours=1)
        Sincronizacion.objects.create(clave='pedidos', actualizado_hasta=inicio)

        def fallar(_filtros: dict):
            for identificador in range(1, 101):
                yield pedido(identificador, 'PAID_WAITING_SHIP')
            raise AnyMarketError('fallo de prueba')

        with patch('modulo.anymarket.services.sincronizacion_pedidos.timezone.now', return_value=ahora), patch(
            'modulo.anymarket.services.sincronizacion_pedidos.consultar_pedidos', side_effect=fallar
        ), self.assertRaises(AnyMarketError):
            sincronizar()
        control = Sincronizacion.objects.get(clave='pedidos')
        self.assertEqual(control.actualizado_hasta, inicio)
        self.assertTrue(control.ultimo_error)
        self.assertEqual(Pedido.objects.count(), 100)
        with patch('modulo.anymarket.services.sincronizacion_pedidos.timezone.now', return_value=ahora), patch(
            'modulo.anymarket.services.sincronizacion_pedidos.consultar_pedidos', return_value=[pedido(i, 'PAID_WAITING_SHIP') for i in range(1, 101)]
        ):
            sincronizar()
        control.refresh_from_db()
        self.assertEqual(Pedido.objects.count(), 100)
        self.assertEqual(control.actualizado_hasta, ahora)
        self.assertEqual(control.ultimo_error, '')
        self.assertEqual(control.ultimo_exito, ahora)

    def test_reconciliacion_actualiza_pedidos_antiguos_sin_avanzar_checkpoint(self) -> None:
        ahora = datetime(2026, 11, 20, tzinfo=dt_timezone.utc)
        previo = ahora - timedelta(hours=2)
        Sincronizacion.objects.create(clave='pedidos', actualizado_hasta=previo)
        Pedido.objects.create(anymarket_id=1, creado_en=ahora, estado='PAID_WAITING_SHIP', datos={})
        with patch('modulo.anymarket.services.sincronizacion_pedidos.timezone.now', return_value=ahora), patch(
            'modulo.anymarket.services.sincronizacion_pedidos.consultar_pedidos', return_value=[pedido(1, 'CANCELED')]
        ) as consulta:
            reconciliar()
        self.assertEqual(consulta.call_count, 2)
        self.assertIn('updatedAfter', consulta.call_args_list[0].args[0])
        self.assertEqual(Pedido.objects.get().estado, 'CANCELED')
        self.assertEqual(Sincronizacion.objects.get(clave='pedidos').actualizado_hasta, previo)
        self.assertEqual(Sincronizacion.objects.get(clave='reconciliacion').ultimo_exito, ahora)

    def test_carga_inicial_reanuda_intervalo_completado(self) -> None:
        ahora = datetime(2027, 5, 1, tzinfo=dt_timezone.utc)
        with patch('modulo.anymarket.services.sincronizacion_pedidos.timezone.now', return_value=ahora), patch(
            'modulo.anymarket.services.sincronizacion_pedidos.consultar_pedidos', side_effect=[[], AnyMarketError('error')]
        ), self.assertRaises(AnyMarketError):
            sincronizar()
        control = Sincronizacion.objects.get(clave='pedidos')
        self.assertIsNone(control.actualizado_hasta)
        self.assertIsNotNone(control.carga_inicial_hasta)
        cursor = control.carga_inicial_hasta
        with patch('modulo.anymarket.services.sincronizacion_pedidos.timezone.now', return_value=ahora), patch(
            'modulo.anymarket.services.sincronizacion_pedidos.consultar_pedidos', return_value=[]
        ) as consulta:
            sincronizar()
        self.assertEqual(consulta.call_args_list[0].args[0]['createdAfter'], cursor.isoformat())
        control.refresh_from_db()
        self.assertEqual(control.actualizado_hasta, ahora)

    @override_settings(ANYMARKET_TOKEN='test', ANYMARKET_PLATFORM='INET', ANYMARKET_ENVIRONMENT='sandbox')
    def test_paginacion_y_cabeceras(self) -> None:
        paginas = [
            {'content': [pedido(1, 'PAID_WAITING_SHIP')], 'links': [{'rel': 'next', 'href': 'https://evil.invalid'}]},
            {'content': [pedido(2, 'PAID_WAITING_SHIP')], 'links': []},
        ]
        with patch('modulo.anymarket.services.sincronizacion_pedidos.build_opener') as opener:
            opener.return_value.open.side_effect = [BytesIO(json.dumps(p).encode()) for p in paginas]
            resultados = list(consultar_pedidos({'status': 'PAID_WAITING_SHIP'}))
        self.assertEqual(len(resultados), 2)
        request = opener.return_value.open.call_args_list[1].args[0]
        self.assertTrue(request.full_url.startswith('https://sandbox-api.anymarket.com.br/v2/orders?'))
        self.assertIn('offset=1', request.full_url)
        self.assertEqual(request.get_header('Gumgatoken'), 'test')

    @override_settings(ANYMARKET_TOKEN='test', ANYMARKET_PLATFORM='INET', ANYMARKET_ENVIRONMENT='sandbox')
    def test_ignora_enlace_next_incorrecto_en_ultima_pagina(self) -> None:
        pagina = {'content': [pedido(1, 'PAID_WAITING_SHIP')],
                  'page': {'number': 1, 'totalPages': 1}, 'links': [{'rel': 'next'}]}
        with patch('modulo.anymarket.services.sincronizacion_pedidos.build_opener') as opener:
            opener.return_value.open.return_value = BytesIO(json.dumps(pagina).encode())
            self.assertEqual(len(list(consultar_pedidos({}))), 1)
            self.assertEqual(opener.return_value.open.call_count, 1)


class BloqueoSincronizacionTests(TransactionTestCase):
    def test_otro_proceso_impide_sincronizacion_y_reconciliacion_reintenta(self) -> None:
        listo, liberar = Event(), Event()

        def mantener_bloqueo() -> None:
            conexion = connections['default']
            try:
                with conexion.cursor() as cursor:
                    cursor.execute('SELECT pg_advisory_lock(%s)', [731984601241])
                    listo.set()
                    liberar.wait(timeout=10)
                    cursor.execute('SELECT pg_advisory_unlock(%s)', [731984601241])
            finally:
                conexion.close()

        hilo = Thread(target=mantener_bloqueo)
        hilo.start()
        try:
            self.assertTrue(listo.wait(timeout=5))
            with patch('modulo.anymarket.services.sincronizacion_pedidos.consultar_pedidos') as consulta:
                sincronizar()
                with self.assertRaises(AnyMarketError):
                    reconciliar()
                consulta.assert_not_called()
        finally:
            liberar.set()
            hilo.join(timeout=10)


@override_settings(STORAGES={
    'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
    'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'},
}, ANYMARKET_TIME_ZONE='America/Santiago')
class PedidosVistaTests(TestCase):
    def setUp(self) -> None:
        self.usuario = get_user_model().objects.create_user(username='operador')
        self.url = '/api/v1/pedidos/'

    def test_requiere_login_y_permiso(self) -> None:
        self.assertEqual(self.client.get(self.url).status_code, 401)
        self.client.force_login(self.usuario)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertFalse(self.client.get('/api/v1/sesion/').json()['puede_ver_pedidos'])

    def test_listado_paginado_escapa_datos_y_muestra_fecha_chilena(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        for identificador in range(26):
            Pedido.objects.create(
                anymarket_id=identificador + 1,
                creado_en=datetime(2026, 10, 27, 4, 0, tzinfo=dt_timezone.utc),
                estado='PAID_WAITING_SHIP',
                datos={'marketPlaceNumber': '<script>malicioso</script>', 'marketPlace': 'TIENDA', 'total': 0,
                       'buyer': {'name': 'Cliente', 'email': 'NO EXPONER'}},
            )
        respuesta = self.client.get(self.url)
        contenido = respuesta.json()
        self.assertEqual(len(contenido['resultados']), 20)
        self.assertEqual(contenido['resultados'][0]['fecha_local'], '27/10/2026 01:00:00')
        self.assertEqual(contenido['resultados'][0]['numero'], '<script>malicioso</script>')
        self.assertNotContains(respuesta, 'NO EXPONER')
        self.assertEqual(len(self.client.get(self.url, {'page': 2}).json()['resultados']), 6)
        self.assertTrue(self.client.get('/api/v1/sesion/').json()['puede_ver_pedidos'])
        self.assertIn('no-store', respuesta['Cache-Control'])

    def test_sin_pedidos(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        self.assertEqual(self.client.get(self.url).json()['resultados'], [])

    def test_estado_sincronizacion_en_api(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        fecha = datetime(2026, 10, 27, 4, tzinfo=dt_timezone.utc)
        Sincronizacion.objects.create(clave='pedidos', ultimo_exito=fecha)
        Sincronizacion.objects.create(clave='reconciliacion', ultimo_error='Error de prueba', ultimo_error_en=fecha)
        resultado = self.client.get(self.url).json()
        self.assertEqual(resultado['zona_horaria'], 'America/Santiago')
        self.assertEqual(resultado['sincronizaciones']['pedidos']['ultimo_exito'], fecha.isoformat())
        self.assertEqual(resultado['sincronizaciones']['reconciliacion']['ultimo_error'], 'Error de prueba')

    def test_estado_fullfillment_en_listado_y_excel(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        for numero, valor in enumerate((True, False, None, 'true'), start=1):
            Pedido.objects.create(anymarket_id=numero, creado_en=datetime(2026, 10, 27, tzinfo=dt_timezone.utc),
                                  estado='PAID_WAITING_SHIP', datos={'fulfillment': valor})
        respuesta = self.client.get(self.url)
        self.assertEqual(sum(p['fulfillment'] for p in respuesta.json()['resultados']), 1)
        exportacion = self.client.get(reverse('anymarket:exportar_pedidos'))
        libro = load_workbook(BytesIO(b''.join(exportacion.streaming_content)), read_only=True)
        valores = [fila[7] for fila in list(libro.active.values)[1:]]
        self.assertEqual(valores.count('FullFillment'), 1)
        self.assertTrue(all(valor in ('FullFillment', None) for valor in valores))

    @patch('modulo.anymarket.views.pedidos.sincronizar_pedidos.apply_async')
    def test_sincronizacion_manual_requiere_permiso_y_limita_solicitudes(self, encolar: MagicMock) -> None:
        cache.clear()
        url = reverse('anymarket:ejecutar_sincronizacion')
        self.assertEqual(self.client.post(url).status_code, 302)
        self.client.force_login(self.usuario)
        self.assertEqual(self.client.post(url).status_code, 403)
        self.usuario.user_permissions.add(Permission.objects.get(codename='change_pedido', content_type__app_label='anymarket'))
        self.client.force_login(get_user_model().objects.get(pk=self.usuario.pk))
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertRedirects(self.client.post(url), '/anymarket/pedidos/', fetch_redirect_response=False)
        self.assertRedirects(self.client.post(url), '/anymarket/pedidos/', fetch_redirect_response=False)
        self.assertRedirects(self.client.get('/anymarket/pedidos/'), '/pedidos/', fetch_redirect_response=False)
        encolar.assert_called_once_with(retry=False)
        cache.clear()

    def test_filtros_fecha_chilena_busqueda_canal_y_rango_invalido(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        for identificador, fecha in enumerate((datetime(2026, 10, 27, 2, 59, tzinfo=dt_timezone.utc),
                                              datetime(2026, 10, 27, 3, 0, tzinfo=dt_timezone.utc)), start=1):
            Pedido.objects.create(anymarket_id=identificador, creado_en=fecha, estado='PAID_WAITING_SHIP',
                                  datos={'marketPlaceNumber': 'OC-123', 'marketPlace': 'Tienda',
                                         'buyer': {'name': 'Ana'}, 'items': [{}, {}], 'total': 100,
                                         'marketPlaceStatus': 'paid'})
        respuesta = self.client.get(self.url, {'desde': '2026-10-27', 'hasta': '2026-10-27',
                                              'busqueda': 'ana', 'canal': 'Tienda', 'estado': 'paid'})
        self.assertEqual(respuesta.json()['total'], 1)
        self.assertEqual(respuesta.json()['resultados'][0]['cliente'], 'Ana')
        self.assertEqual(respuesta.json()['resultados'][0]['items'], 2)
        self.assertEqual(self.client.get(self.url, {'desde': '2026-10-28', 'hasta': '2026-10-27'}).status_code, 400)

    def test_orden_total_across_paginas_y_parametros_invalidos(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        for numero in range(26):
            Pedido.objects.create(anymarket_id=numero + 1, creado_en=datetime(2026, 10, 27, tzinfo=dt_timezone.utc),
                                  estado='PAID_WAITING_SHIP', datos={'total': numero})
        primera = self.client.get(self.url, {'orden': 'total', 'direccion': 'desc'})
        self.assertEqual(primera.json()['resultados'][0]['total'], 25)
        segunda = self.client.get(self.url, {'orden': 'total', 'direccion': 'desc', 'page': 2})
        self.assertEqual(segunda.json()['resultados'][-1]['total'], 0)
        invalida = self.client.get(self.url, {'orden': 'datos__buyer', 'direccion': 'invalida'})
        self.assertEqual(invalida.status_code, 200)

    def test_busqueda_global_abarca_columnas_y_todas_las_paginas(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        fecha = datetime(2026, 10, 27, 4, 0, tzinfo=dt_timezone.utc)
        Pedido.objects.create(anymarket_id=1, creado_en=fecha, estado='PAID_WAITING_SHIP', datos={
            'marketPlaceNumber': 'OC-ESPECIAL', 'marketPlace': 'CANAL-UNICO',
            'buyer': {'name': 'CLIENTE-UNICO'}, 'items': [{}, {}, {}],
            'total': 98765, 'marketPlaceStatus': 'ESTADO-UNICO',
        })
        for numero in range(2, 15):
            Pedido.objects.create(anymarket_id=numero, creado_en=fecha, estado='PAID_WAITING_SHIP', datos={})
        for termino in ('OC-ESPECIAL', 'CANAL-UNICO', 'CLIENTE-UNICO', 'ESTADO-UNICO', '98765', '3'):
            with self.subTest(termino=termino):
                respuesta = self.client.get(self.url, {'busqueda': termino, 'tamano': '5'})
                self.assertEqual(respuesta.json()['total'], 1)
                self.assertEqual(respuesta.json()['resultados'][0]['numero'], 'OC-ESPECIAL')
        respuesta = self.client.get(self.url, {'busqueda': '27/10/2026'})
        self.assertEqual(respuesta.json()['total'], 14)
        exportacion = self.client.get(reverse('anymarket:exportar_pedidos'), {'busqueda': 'CANAL-UNICO'})
        libro = load_workbook(BytesIO(b''.join(exportacion.streaming_content)), read_only=True)
        self.assertEqual(len(list(libro.active.rows)), 2)

    def test_nombres_de_canal_en_listado_busqueda_y_exportacion(self) -> None:
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(self.usuario)
        for numero, (original, nombre) in enumerate((('mercado_livre', 'MERCADO LIBRE'),
                                                     ('SHOPIFY', 'PORTOMENAJE'),
                                                     ('C52434-09', 'YAKU')), start=1):
            with self.subTest(canal=original):
                Pedido.objects.create(anymarket_id=numero, creado_en=datetime(2026, 10, 27, tzinfo=dt_timezone.utc),
                                      estado='PAID_WAITING_SHIP', datos={'marketPlace': original})
                respuesta = self.client.get(self.url, {'busqueda': nombre, 'canal': original})
                self.assertEqual(respuesta.json()['total'], 1)
                self.assertEqual(respuesta.json()['resultados'][0]['canal'], original)
                exportacion = self.client.get(reverse('anymarket:exportar_pedidos'), {'busqueda': nombre})
                libro = load_workbook(BytesIO(b''.join(exportacion.streaming_content)), read_only=True)
                self.assertEqual(list(libro.active.values)[1][0], nombre)

    def test_exportacion_filtrada_protegida_y_sin_formulas(self) -> None:
        url = reverse('anymarket:exportar_pedidos')
        self.assertEqual(self.client.get(url).status_code, 302)
        self.client.force_login(self.usuario)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.usuario.user_permissions.add(Permission.objects.get(codename='view_pedido', content_type__app_label='anymarket'))
        self.client.force_login(get_user_model().objects.get(pk=self.usuario.pk))
        for numero in range(12):
            Pedido.objects.create(anymarket_id=numero + 1, creado_en=datetime(2026, 10, 27, 4, tzinfo=dt_timezone.utc),
                                  estado='PAID_WAITING_SHIP', datos={'marketPlace': 'TIENDA', 'marketPlaceStatus': 'paid',
                                  'marketPlaceNumber': '=1+1', 'buyer': {'name': 'Cliente'}, 'total': numero})
        Pedido.objects.create(anymarket_id=100, creado_en=datetime(2026, 10, 27, tzinfo=dt_timezone.utc),
                              estado='PAID_WAITING_SHIP', datos={'marketPlace': 'OTRA', 'marketPlaceStatus': 'pending'})
        respuesta = self.client.get(url, {'canal': 'TIENDA', 'estado': 'paid', 'tamano': '5'})
        self.assertEqual(respuesta.status_code, 200)
        libro = load_workbook(BytesIO(b''.join(respuesta.streaming_content)), read_only=True)
        filas = list(libro.active.rows)
        self.assertEqual(len(filas), 13)
        self.assertEqual(filas[1][1].value, '=1+1')
        self.assertEqual(filas[1][1].data_type, 's')
        self.assertEqual(filas[1][6].value, 'paid')
        self.assertEqual(self.client.get(url, {'desde': 'invalida'}).status_code, 400)
